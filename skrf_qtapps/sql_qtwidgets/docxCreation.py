import os
import docx
import tempfile

import pyodbc

from skrf import Network

from docx.enum.text import WD_BREAK

from copy import deepcopy

import numpy as np
import matplotlib.pyplot as plt

from datetime import date, datetime

file_types = ('.csv', '.s1p', '.s2p', '.s3p', '.s4p')
testDataPath = '\\\\Filesrv\\Test\\RFData\\'

gang_plots_dict = {'Return Loss': ['return loss', 'rl', 's11', 's22', 'return_loss'],
                   'Insertion Loss': ['insertion loss', 'il', 'insertion_loss', 'transmission'],
                   'Attenuation': ['attenuation'], 
                   'Isolation': ['isolation', 'isol']}


plt.rcParams["font.family"] = "Calibri"

def create_docx_report(instrument_dictionary, revision_number='1', selected_file_only=[], gang_plots=False):

    date_now = date.today().strftime("%d/%m/%Y")

    for partID in instrument_dictionary:

        # Get Part ID Information
        mydb = pyodbc.connect("DRIVER={SQL Server};SERVER=SQLSRV22;DATABASE=ISM;UID=FLUser;PWD=MelonBall", readonly=True)
        mydb_cursor = mydb.cursor()
        partid_sql_info = mydb_cursor.execute("select Instrument_Number, Instrument_ID, Part_ID, Series, Var_Suffix, var_id, SeriesName from vw_Instrument_VarDetails where (Part_ID = ?)",(partID)).fetchone()
        mydb.close()

        instrument_number = partid_sql_info.Instrument_Number
        instrument_id = str(partid_sql_info.Instrument_ID)
        var_suffix = partid_sql_info.Var_Suffix
        seriesName = partid_sql_info.SeriesName

        if var_suffix is None:
            var_suffix = ''

        model = f"{instrument_number} {var_suffix}"
        
        for sn in instrument_dictionary[partID]:
            # Open Template
            doc = docx.Document(os.path.abspath(os.path.join(os.path.dirname(__file__), "report_template.docx")))

            # Table Data
            table = doc.tables[0]

            paragraph = table.cell(0, 0).add_paragraph()
            run = paragraph.add_run(sn)
            run.font.size = docx.shared.Pt(9)
            run.font.name = 'Calibri'
            run.italic = False

            paragraph = table.cell(0, 1).add_paragraph()
            run = paragraph.add_run(model)
            run.font.size = docx.shared.Pt(9)
            run.font.name = 'Calibri'
            run.italic = False

            paragraph = table.cell(0, 2).add_paragraph()
            run = paragraph.add_run(partID)
            run.font.size = docx.shared.Pt(9)
            run.font.name = 'Calibri'
            run.italic = False

            paragraph = table.cell(0, 3).add_paragraph()
            run = paragraph.add_run(seriesName)
            run.font.size = docx.shared.Pt(9)
            run.font.name = 'Calibri'
            run.italic = False

            paragraph = table.cell(1, 2).add_paragraph()
            run = paragraph.add_run(date_now)
            run.font.size = docx.shared.Pt(9)
            run.font.name = 'Calibri'
            run.italic = False

            paragraph = table.cell(1, 3).add_paragraph()
            run = paragraph.add_run(instrument_id)
            run.font.size = docx.shared.Pt(9)
            run.font.name = 'Calibri'
            run.italic = False

            paragraph = table.cell(1, 4).add_paragraph()
            run = paragraph.add_run(revision_number)
            run.font.size = docx.shared.Pt(9)
            run.font.name = 'Calibri'
            run.italic = False

            # Add Plots
            # try:
            file_list = os.listdir(testDataPath+partID)
            file_list = [name for name in file_list if name.lower().endswith(file_types)]
            file_list = [name for name in file_list if sn in name]
            # list_files = list(filter(lambda f: not any(s in f.lower() for s in ['top_of_band', 'bottom_of_band']), list_files))

            if selected_file_only:
                file_list = [f for f in file_list if f in selected_file_only]

            new_files_list = []

            if gang_plots:
                for key in gang_plots_dict:
                    filter_list = gang_plots_dict[key]
                    gang_files = [name for name in file_list if any(s in name.lower() for s in filter_list)]
                    if len(gang_files) == 1:  # Only create gang plot if more than 1 relevant file is found
                        continue
                    file_list = list(set(file_list) - set(gang_files))  # Remove gang plot files from individual plotting
                    if gang_files:
                        new_files_list.append(gang_files + [key])  # Append gang plot files as a group at the end of the list

            file_list = file_list + new_files_list  # Append gang plot groups at the end of the list

            file_counter = 0
            length_file_list = len(file_list)

            for file_list_element in file_list:
                # Duplicating table for each file after the first
                if file_list.index(file_list_element) != 0:
                    orig_table = doc.tables[0]

                    tbl_element = deepcopy(orig_table._element)
                    new_table = doc.add_table(rows=1, cols=1)  # Temporary table to get a reference
                    new_table._element.addnext(tbl_element)

                    doc.add_paragraph()  # Add a paragraph break after the table
                
                if isinstance(file_list_element, list) or file_list_element.lower().endswith('.csv'):
                    '''Original CSV Format'''
                    if file_list.index(file_list_element) == 0:
                        if isinstance(file_list_element, list):
                            parameter_file = file_list_element[0]  # Get the first file name from the gang plot group for metadata extraction
                        else:
                            parameter_file = file_list_element
                        # Get Analyser and Tester from CSV
                        parameters = np.loadtxt(testDataPath+partID+'\\'+parameter_file, delimiter=',', usecols=2, dtype=str)
                        analyser = parameters[2][8:]
                        tester = parameters[0][2:]

                        paragraph = table.cell(1, 0).add_paragraph()
                        run = paragraph.add_run(analyser)
                        run.font.size = docx.shared.Pt(9)
                        run.font.name = 'Calibri'
                        run.italic = False

                        paragraph = table.cell(1, 1).add_paragraph()
                        run = paragraph.add_run(tester)
                        run.font.size = docx.shared.Pt(9)
                        run.font.name = 'Calibri'
                        run.italic = False

                    if isinstance(file_list_element, str):
                        title = file_list_element[:-4].replace('_', ' ')
                        file_list_element = [file_list_element]  # Convert to list for consistent processing in gang and individual plots

                    elif isinstance(file_list_element, list):
                        title = f"Gang Plot - {file_list_element[-1]}"  # Last element in the list is the gang plot type (e.g., 'Return Loss')
                        file_list_element = file_list_element[:-1]  # Remove the gang plot type from the list to get the actual file names
                    
                    fig, ax = plt.subplots(figsize=(8, 10))

                    for f in file_list_element:

                        data = np.loadtxt(testDataPath+partID+'\\'+f, delimiter=',', usecols=(0, 1, 3))
                        freq = data[:, 0] 
                        mag_s_param = data[:, 1]
                        tested_specification = data[:, 2]

                        legend_name = f[:-4].replace('_', ' ')
                        
                        if min(mag_s_param) >= min(tested_specification):
                            y_min = 2*min(mag_s_param)
                        else:
                            y_min = max([min(mag_s_param), 2*min(tested_specification), -100])

                        ax.plot(freq, mag_s_param, label=legend_name)

                    ax.plot(freq, tested_specification, '--', label='Specification', color='#004124')
                    ax.set_xlabel('Frequency (GHz)')
                    ax.set_ylabel('Magnitude (dB)')
                    ax.grid(True, which='both')
                    ax.minorticks_on()
                    ax.grid(which='minor', linewidth=0.5, alpha=0.5)
                    ax.set_ylim(top=0, bottom=y_min)
                    if gang_plots:
                        ax.legend(fontsize='small')
                    ax.set_title(title)
                    fig.tight_layout()

                    tempdir = tempfile.gettempdir()
                    img_path = os.path.join(tempdir, f"test_data_print_temp.png")
                    plt.savefig(img_path)
                    plt.close()
                    paragraph = doc.add_paragraph()
                    run = paragraph.add_run()
                    run.add_picture(img_path, width=docx.shared.Inches(6))
                    paragraph.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
                    os.remove(img_path)
                
                elif file_list_element.lower().endswith(('.s1p', '.s2p', '.s3p', '.s4p')):
                    '''SNP Format'''
                    ntwk = Network(testDataPath+partID+'\\'+file_list_element)
                    fig = ntwk.plot_s_db(m=0, x='freq', show_legend=False)
                    plt.title(f"{file_list_element}")

                    tempdir = tempfile.gettempdir()
                    img_path = os.path.join(tempdir, f"test_data_print_temp.png")
                    plt.savefig(img_path)
                    plt.close()

                    doc.add_picture(img_path, width=docx.shared.Inches(6))
                    
                    os.remove(img_path)

                if file_counter < length_file_list - 1:  # Add page break after each file except the last one
                    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
                    file_counter += 1
            # except:
            #     print(f"Could not find directory for PartID: {partID}")
            #     continue
            
            doc.core_properties.last_modified_by = os.getlogin()
            doc.core_properties.modified = datetime.now()
            doc.core_properties.comments = f"Test Report for {sn} - Revision {revision_number}"
            doc.core_properties.revision = int(revision_number)
            doc.save(f"{sn}_test_report_R{revision_number}.docx")#{testDataPath+partID+'\\'

create_docx_report({'F01999': ['278751'], 'F06416': ['323379']}, gang_plots=True)